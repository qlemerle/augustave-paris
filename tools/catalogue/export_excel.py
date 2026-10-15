"""Crée le fichier Excel des bijoux (une ligne par bijou) à partir du catalogue.

Usage : python3 tools/catalogue/export_excel.py [sortie.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

from common import COLUMNS, DEFAULT_SIZES, HEADERS, ROOT, TAGS, load

out = sys.argv[1] if len(sys.argv) > 1 else "augustave-bijoux.xlsx"
data = load()
col_title = {c["handle"]: c["title"] for c in data["collections"]}
cat_title = {c["handle"]: c["title"] for c in data["categories"]}
photos_dir = ROOT / "public/images/produits"

NAVY, GOLD, CREAM = "0E2340", "C6A15B", "F4F2EF"
wb = Workbook()
ws = wb.active
ws.title = "Bijoux"
ws.append(HEADERS)
for i, (name, width, help_) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=i)
    cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.comment = Comment(help_, "Augustave")
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 34
ws.freeze_panes = "C2"

for p in data["products"]:
    sizes = p.get("sizes") or DEFAULT_SIZES.get(p["options"], [])
    photos = sorted(f.name for f in photos_dir.glob(p["handle"] + "-*")) if photos_dir.exists() else []
    ws.append([
        p["handle"],
        p["title"],
        col_title.get(p["collection"], p["collection"]),
        cat_title.get(p["category"], p["category"]),
        p["price"],
        p.get("compareAtPrice"),
        " / ".join(sizes),
        *["Oui" if tag in p["tags"] else None for _, tag in TAGS],
        p.get("excerpt", ""),
        p.get("description", ""),
        p.get("materials", ""),
        p.get("dimensions"),
        ", ".join(photos) or f"{p['handle']}-1.jpg, {p['handle']}-2.jpg",
    ])

last = ws.max_row + 200  # place pour de nouveaux bijoux
thin = Side(style="thin", color="E2DED8")
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    row[0].font = Font(color="8A8F96")
    row[-1].font = Font(color="8A8F96", italic=True)
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 60
for col in ("E", "F"):
    for r in range(2, last):
        ws[f"{col}{r}"].number_format = '#,##0.00 "€"'

# Listes déroulantes
lists = wb.create_sheet("Listes")
lists.append(["Collections", "Catégories", "Oui"])
cols = [c["title"] for c in data["collections"]]
cats = [c["title"] for c in data["categories"]]
for i in range(max(len(cols), len(cats))):
    lists.append([cols[i] if i < len(cols) else None, cats[i] if i < len(cats) else None, "Oui" if i == 0 else None])
dv_col = DataValidation(type="list", formula1=f"=Listes!$A$2:$A${len(cols) + 1}", allow_blank=False)
dv_cat = DataValidation(type="list", formula1=f"=Listes!$B$2:$B${len(cats) + 1}", allow_blank=False)
dv_yes = DataValidation(type="list", formula1='"Oui"', allow_blank=True)
dv_price = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
for dv in (dv_col, dv_cat, dv_yes, dv_price):
    ws.add_data_validation(dv)
dv_col.add(f"C2:C{last}")
dv_cat.add(f"D2:D{last}")
dv_price.add(f"E2:F{last}")
first_tag = 8
from openpyxl.utils import get_column_letter
dv_yes.add(f"{get_column_letter(first_tag)}2:{get_column_letter(first_tag + len(TAGS) - 1)}{last}")
lists.sheet_state = "hidden"

# Mode d'emploi
help_ws = wb.create_sheet("Mode d'emploi", 0)
lines = [
    ("Fichier des bijoux Augustave Paris", True),
    ("", False),
    ("Une ligne par bijou, dans l'onglet « Bijoux ». Modifiez librement les prix et les textes.", False),
    ("Ne changez pas la colonne « Référence » : c'est elle qui relie la ligne au bijou du site.", False),
    ("Pour ajouter un bijou : remplissez une nouvelle ligne en laissant la référence vide.", False),
    ("Pour retirer un bijou : supprimez sa ligne.", False),
    ("Tailles proposées : séparées par « / » (ex. 40 cm / 45 cm). Vide = tailles habituelles de la catégorie.", False),
    ("Pièce unique, Nouveauté, Coup de cœur, Édition limitée : mettez « Oui » pour afficher le badge.", False),
    ("Photos : déposez-les sur le Drive en les nommant comme indiqué dans la dernière colonne (ex. collier-rouge-love-1.jpg).", False),
    ("", False),
    ("Renvoyez ensuite le fichier : le site est mis à jour à partir de celui-ci.", False),
]
for text, bold in lines:
    help_ws.append([text])
    c = help_ws.cell(row=help_ws.max_row, column=1)
    c.font = Font(bold=bold, size=16 if bold else 12, color=NAVY)
help_ws.column_dimensions["A"].width = 120
help_ws.sheet_view.showGridLines = False
wb.active = 1

wb.save(out)
print(f"{out} : {len(data['products'])} bijoux")

"""Met à jour le catalogue (lib/commerce/mock/catalog.json) depuis le fichier Excel des bijoux.

Usage : python3 tools/catalogue/import_excel.py augustave-bijoux.xlsx [--dry-run]
"""
import re
import sys
import unicodedata

from openpyxl import load_workbook

from common import DEFAULT_SIZES, HEADERS, OPTION_BY_CATEGORY, TAGS, load, save

path = sys.argv[1]
dry = "--dry-run" in sys.argv
data = load()
by_handle = {p["handle"]: p for p in data["products"]}
col_by_title = {c["title"].lower(): c["handle"] for c in data["collections"]}
col_by_title.update({c["handle"]: c["handle"] for c in data["collections"]})
cat_by_title = {c["title"].lower(): c["handle"] for c in data["categories"]}
cat_by_title.update({c["handle"]: c["handle"] for c in data["categories"]})


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def text(v):
    return re.sub(r"\s+\n", "\n", str(v).strip()) if v is not None else ""


def money(v, label, row):
    if v in (None, ""):
        return None
    try:
        n = float(str(v).replace("€", "").replace(",", ".").replace(" ", "").strip())
    except ValueError:
        sys.exit(f"Ligne {row} : « {label} » n'est pas un prix ({v!r}).")
    return int(n) if n.is_integer() else round(n, 2)


ws = load_workbook(path, data_only=True)["Bijoux"]
head = [text(c.value) for c in ws[1]]
missing = [h for h in HEADERS if h not in head]
if missing:
    sys.exit(f"Colonnes introuvables : {', '.join(missing)}")
idx = {h: head.index(h) for h in HEADERS}

products, report, seen = [], [], set()
for row in ws.iter_rows(min_row=2, values_only=True):
    get = lambda h: row[idx[h]] if idx[h] < len(row) else None
    name = text(get("Nom"))
    if not name:
        continue
    r = len(products) + 2
    handle = text(get("Référence")) or slug(name)
    if handle in seen:
        sys.exit(f"Ligne {r} : référence en double ({handle}).")
    seen.add(handle)
    collection = col_by_title.get(text(get("Collection")).lower())
    category = cat_by_title.get(text(get("Catégorie")).lower())
    if not collection:
        sys.exit(f"Ligne {r} ({name}) : collection inconnue « {get('Collection')} ».")
    if not category:
        sys.exit(f"Ligne {r} ({name}) : catégorie inconnue « {get('Catégorie')} ».")
    price = money(get("Prix (€)"), "Prix", r)
    if price is None:
        sys.exit(f"Ligne {r} ({name}) : prix manquant.")
    old = by_handle.get(handle)
    p = dict(old) if old else {}
    if not old:
        # nouveau bijou : illustration de remplacement reprise d'un bijou de la même collection
        twin = next((x for x in data["products"] if x["collection"] == collection), data["products"][0])
        p["art"] = twin["art"]
        p["options"] = OPTION_BY_CATEGORY.get(category, "unique")
    p.update(handle=handle, title=name, collection=collection, category=category, price=price)
    p["tags"] = [tag for label, tag in TAGS if text(get(label)).lower() in ("oui", "x", "yes", "o")]
    p["excerpt"], p["description"], p["materials"] = text(get("Accroche")), text(get("Description")), text(get("Matières"))
    for key, value in (("compareAtPrice", money(get("Prix barré (€)"), "Prix barré", r)), ("dimensions", text(get("Dimensions")) or None)):
        if value is None:
            p.pop(key, None)
        else:
            p[key] = value
    sizes = [s.strip() for s in text(get("Tailles proposées")).split("/") if s.strip()]
    if sizes and sizes != DEFAULT_SIZES.get(p["options"]):
        p["sizes"] = sizes
    else:
        p.pop("sizes", None)
    # ordre des clés stable
    order = ["handle", "title", "collection", "category", "price", "compareAtPrice", "tags", "excerpt", "description", "materials", "dimensions", "options", "sizes", "art"]
    p = {k: p[k] for k in order if k in p}
    if not old:
        report.append(f"+ {handle} (nouveau)")
    else:
        changes = [k for k in p if k != "art" and old.get(k) != p.get(k)] + [k for k in old if k not in p]
        if changes:
            report.append(f"~ {handle} : {', '.join(changes)}" + (f" ({old['price']} € → {p['price']} €)" if "price" in changes else ""))
    products.append(p)

for h in by_handle:
    if h not in seen:
        report.append(f"- {h} (retiré)")

print("\n".join(report) or "Aucun changement.")
print(f"{len(products)} bijoux dans le fichier.")
if not dry:
    data["products"] = products
    save(data)
    print("Catalogue mis à jour.")
